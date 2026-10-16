"""Reproduce the revised manuscript analyses from the archived original ratings.

No model/API calls, score transformations, or changes to source files.
See README.md for the data layout, bootstrap assumptions and output files.
"""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import platform
import re
import numpy as np
import openpyxl
import scipy
from scipy import stats

MODELS = ["ChatGPT", "Deepseek", "Llama", "Claude", "Gemini"]

def extract_scores(raw):
    lines = raw.replace('*', '').replace('`', '').splitlines()
    scores = {}
    risk_column = None
    for line in lines:
        if not line.strip().startswith('|'):
            continue
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        columns = [i for i, c in enumerate(cells)
                   if re.search(r'(?i)risk.*(?:score|0|10)|risk score', c)]
        if columns:
            risk_column = columns[0]
            continue
        if (risk_column is not None and cells[0].isdigit()
                and 1 <= int(cells[0]) <= 18 and len(cells) > risk_column):
            value = re.match(r'(\d+(?:\.\d+)?)', cells[risk_column])
            if value:
                scores[int(cells[0])] = float(value[1])
    if len(scores) == 18:
        return scores
    values = []
    for line in lines:
        if re.search(r'(?i)overall|average|summary|primary concern|secondary concern', line):
            continue
        match = re.search(r'(?i)risk\s*(?:level|rating|score)?\s*:?\s*(\d+(?:\.\d+)?)', line)
        if match:
            values.append(float(match[1]))
            if len(values) == 18:
                break  # Subsequent overall ratings are outside the item-level analysis.
    if len(values) != 18:
        raise ValueError(f'Expected 18 item scores, found {len(values)}')
    return dict(enumerate(values, 1))


def interval(values):
    return [float(v) for v in np.quantile(values, [.025, .975])]


def holm(pvalues):
    order = np.argsort(pvalues)
    adjusted = np.empty(len(pvalues))
    adjusted[order] = np.minimum(1, np.maximum.accumulate(
        np.array(pvalues)[order] * (len(pvalues) - np.arange(len(pvalues)))))
    return adjusted


def metrics(x, y):
    difference = x - y
    test = stats.ttest_rel(x, y)
    ci = test.confidence_interval()
    return dict(mean_llm=float(x.mean()), mean_human=float(y.mean()),
                bias=float(difference.mean()), mae=float(abs(difference).mean()),
                rmse=float(np.sqrt(np.mean(difference ** 2))),
                pearson=float(stats.pearsonr(x, y).statistic),
                # Suppress floating-point tie splitting at decimal score precision.
                spearman=float(stats.spearmanr(np.round(x, 12), np.round(y, 12)).statistic),
                dz=float(difference.mean() / difference.std(ddof=1)),
                t=float(test.statistic), df=int(test.df), p=float(test.pvalue),
                t_ci_bias=[float(ci.low), float(ci.high)])


def write_csv(path, rows):
    with path.open('w', newline='', encoding='utf-8') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def analyze(data_dir, output_dir, resamples=20000, seed=20261005):
    source = data_dir / 'results_final_scenarios.xlsx'
    human_source = data_dir / 'pesquisa-de-riscos-ciberneticos-respostas.csv'
    inputs = [source.resolve(), human_source.resolve()]
    workbook = openpyxl.load_workbook(source, data_only=True)
    comparison = workbook['Comparações']
    # Axes: model, item, conversation. Ignore cached rounded means.
    scores = np.array([[[comparison.cell(i, c).value
                         for c in range(start, start + 5)]
                        for i in range(3, 21)]
                       for start in [2, 8, 14, 20, 26]], dtype=float)
    audit = []
    for m, model in enumerate(MODELS):
        for run in range(5):
            raw = workbook[model].cell(3, run + 1).value
            if not isinstance(raw, str):
                raise ValueError(f'Missing archived response: {model}, run {run + 1}')
            for item, value in extract_scores(raw).items():
                archived = float(scores[m, item - 1, run])
                audit.append(dict(model=model, run=run + 1, item=item,
                                  text_score=value, transcribed_score=archived,
                                  match=bool(value == archived)))
    if len(audit) != 450 or not all(row['match'] for row in audit):
        raise ValueError('Model response text does not match all 450 transcribed ratings.')

    with human_source.open(encoding='utf-8-sig', newline='') as stream:
        raw_rows = list(csv.reader(stream, delimiter=';'))
    rows = raw_rows[1:]
    if len(rows) != 50 or any(len(row) != 41 for row in rows):
        raise ValueError('Expected 50 human records with 41 CSV columns each.')
    # Column 3: self-reported professional level; 5,7,...,39: Q1--Q18 ratings.
    humans = np.array([[float(row[i]) for i in range(5, 41, 2)] for row in rows])
    senior = np.array([row[3] in ['Sênior', 'Especialista'] for row in rows])
    if senior.sum() != 32 or (~senior).sum() != 18:
        raise ValueError('Expected 32 senior/specialist participants and 18 others.')
    for name, array in [('Human', humans), ('Model', scores)]:
        if not np.isfinite(array).all() or not ((array >= 0) & (array <= 10)).all():
            raise ValueError(f'{name} ratings must be finite and within 0--10.')
    provenance = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs}

    # PCG64 is the generator used by default_rng in the original analysis.
    # Preserve drawing order: Group 1, Group 2, then each model in MODELS order.
    rng = np.random.Generator(np.random.PCG64(seed))
    strata = [humans[senior], humans[~senior]]
    human_boots_by_stratum = [
        group[rng.integers(0, len(group), (resamples, len(group)))].mean(axis=1)
        for group in strata
    ]
    human_means = {'All': humans.mean(0),
                   'Senior': strata[0].mean(0), 'Other': strata[1].mean(0)}
    human_boots = {
        'All': (32 * human_boots_by_stratum[0] +
                18 * human_boots_by_stratum[1]) / 50,
        'Senior': human_boots_by_stratum[0],
        'Other': human_boots_by_stratum[1],
    }
    # Resample whole 18-item conversation profiles, not isolated item scores.
    model_boots = [array.T[rng.integers(0, 5, (resamples, 5))].mean(axis=1)
                   for array in scores]
    results = []
    for group, human_mean in human_means.items():
        for m, model in enumerate(MODELS):
            values = metrics(scores[m].mean(axis=1), human_mean)
            differences = model_boots[m] - human_boots[group]
            values.update(
                bootstrap_ci_bias=interval(differences.mean(1)),
                bootstrap_ci_mae=interval(abs(differences).mean(1)),
                bootstrap_ci_rmse=interval(np.sqrt((differences ** 2).mean(1))),
                treatment='original', group=group, model=model)
            results.append(values)
    # Family of 15 tests, not the 30-test exploratory family in early drafts.
    for row, adjusted in zip(results, holm([row['p'] for row in results])):
        row['p_holm_15'] = float(adjusted)

    participant_means = [array.mean(1) for array in strata]
    subgroup_difference = human_means['Senior'] - human_means['Other']
    welch = stats.ttest_ind(*participant_means, equal_var=False)
    human_comparison = dict(
        mean_senior=float(participant_means[0].mean()),
        mean_other=float(participant_means[1].mean()),
        difference=float(subgroup_difference.mean()),
        bootstrap_ci_difference=interval(
            (human_boots_by_stratum[0] - human_boots_by_stratum[1]).mean(1)),
        pearson_item_means=float(stats.pearsonr(
            human_means['Senior'], human_means['Other']).statistic),
        mae_item_means=float(abs(subgroup_difference).mean()),
        rmse_item_means=float(np.sqrt(np.mean(subgroup_difference ** 2))),
        welch_t=float(welch.statistic), welch_p=float(welch.pvalue))
    variability, disagreement, item_rows, correlation_rows = {}, {}, [], []
    for group, array in [('All', humans), ('Senior', strata[0]), ('Other', strata[1])]:
        pairs = np.triu_indices(len(array), k=1)
        within = float(abs(array[pairs[0]] - array[pairs[1]]).mean())
        variability[group] = dict(mean_item_sd=float(array.std(0, ddof=1).mean()),
                                  mean_pairwise_absolute_difference=within)
        disagreement[group] = dict(within_human=within, human_model={
            model: float(abs(array[:, None, :] - scores[m].T[None, :, :]).mean())
            for m, model in enumerate(MODELS)})
        for j in range(18):
            item_rows.append(dict(group=group, item=j + 1,
                                  human_mean=float(array[:, j].mean()),
                                  human_sd=float(array[:, j].std(ddof=1)),
                                  **{model: float(scores[m, j].mean())
                                     for m, model in enumerate(MODELS)}))
        average_model_profile = scores.mean(axis=2).mean(axis=0)
        correlation_rows.append(dict(
            group=group, mean_llm=float(average_model_profile.mean()),
            mean_human=float(array.mean()),
            pearson_average_profile=float(stats.pearsonr(
                average_model_profile, array.mean(0)).statistic)))
    disagreement['between_human_groups'] = float(
        abs(strata[0][:, None, :] - strata[1][None, :, :]).mean())
    for m, model in enumerate(MODELS):
        deviations = scores[m].std(1, ddof=1)
        variability[model] = dict(mean_item_run_sd=float(deviations.mean()),
                                 min_item_run_sd=float(deviations.min()),
                                 max_item_run_sd=float(deviations.max()),
                                 identical_items=int((np.ptp(scores[m], axis=1) == 0).sum()))
    output = dict(seed=seed, bootstrap_resamples=resamples, rng='PCG64',
                  percentile_method='linear', source_sha256=provenance,
                  versions={'python': platform.python_version(), 'numpy': np.__version__,
                            'scipy': scipy.__version__, 'openpyxl': openpyxl.__version__},
                  transcription_matches=len(audit), results=results,
                  human_comparison=human_comparison, variability=variability,
                  individual_disagreement=disagreement,
                  average_profile_comparisons=correlation_rows)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_paths = ['results.json', 'metrics.csv', 'transcription_audit.csv',
                    'item_summary.csv', 'average_profile_comparisons.csv']
    if any((output_dir / name).resolve() in inputs for name in output_paths):
        raise ValueError('An output file would overwrite an input.')
    (output_dir / 'results.json').write_text(
        json.dumps(output, indent=2, ensure_ascii=False, allow_nan=False), encoding='utf-8')
    write_csv(output_dir / 'metrics.csv', results)
    write_csv(output_dir / 'transcription_audit.csv', audit)
    write_csv(output_dir / 'item_summary.csv', item_rows)
    write_csv(output_dir / 'average_profile_comparisons.csv', correlation_rows)
    for path in inputs:
        if hashlib.sha256(path.read_bytes()).hexdigest() != provenance[path.name]:
            raise RuntimeError(f'Input changed during analysis: {path.name}')
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data-dir', type=Path,
                        default=Path(__file__).resolve().parent.parent,
                        help='Directory containing the two archived data files (default: repository root).')
    parser.add_argument('--output-dir', type=Path,
                        default=Path(__file__).resolve().parent / 'results')
    parser.add_argument('--resamples', type=int, default=20000)
    parser.add_argument('--seed', type=int, default=20261005)
    args = parser.parse_args()
    if args.resamples < 1 or args.seed < 0:
        parser.error('resamples must be positive and seed must be nonnegative.')
    output = analyze(args.data_dir.resolve(), args.output_dir.resolve(),
                     args.resamples, args.seed)
    print(f"Checked {output['transcription_matches']} archived scores; "
          f"computed {len(output['results'])} original-scale comparisons.")
    print(f"Outputs: {args.output_dir.resolve()}")


if __name__ == '__main__':
    main()
